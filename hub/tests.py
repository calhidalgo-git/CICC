import json
import os
import re
import subprocess
import sys
import tempfile
from datetime import timedelta
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import Group, User
from django.contrib.sessions.models import Session
from django.core import mail, serializers
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from .catalog import CONTENT, MODULES
from .models import (
    Assessment,
    Attempt,
    Audit,
    Certificate,
    CourseWeek,
    Favorite,
    LearningPath,
    Lesson,
    Notification,
    Progress,
    Question,
    Resource,
    ResourceVersion,
    WeekProgress,
)
from .navigation import NAV
from .services import draw_questions, is_correct


def valid_template(client):
    """Plantilla de importación con el ID real de la primera evaluación (los IDs no siempre empiezan en 1)."""
    wb = load_workbook(BytesIO(client.get('/template/').content))
    wb.active['A2'] = Assessment.objects.order_by('pk').first().pk
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


PASSWORD = 'DemoQA!2026'


@override_settings(
    MEDIA_ROOT=tempfile.mkdtemp(prefix='cicc-test-media-'),
    PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'],
)
class BaseCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('seed_demo', password=PASSWORD)

    def setUp(self):
        cache.clear()

    def login(self, name='ana'):
        self.assertTrue(self.client.login(username=name, password=PASSWORD))

    def start_attempt(self, assessment=None):
        self.login()
        assessment = assessment or Assessment.objects.first()
        self.client.post(f'/exam/{assessment.pk}/start/')
        return Attempt.objects.latest('pk')


class AccessAndScreens(BaseCase):
    def test_auth_and_roles(self):
        self.assertEqual(self.client.get('/').status_code, 302)
        self.login()
        self.assertEqual(self.client.get('/p/users/').status_code, 403)
        self.assertEqual(self.client.get('/report.xlsx').status_code, 403)
        self.assertEqual(self.client.get('/p/questions/').status_code, 403)

    def test_all_admin_screens(self):
        self.login('admin')
        for slug, _, _ in NAV:
            self.assertEqual(self.client.get(f'/p/{slug}/').status_code, 200, slug)
        resource = Resource.objects.first()
        self.assertEqual(self.client.get(f'/resource/{resource.pk}/').status_code, 200)

    def test_unknown_section_is_404(self):
        self.login()
        self.assertEqual(self.client.get('/p/nope/').status_code, 404)

    def test_modules_and_brand(self):
        self.login()
        for slug in MODULES:
            response = self.client.get(f'/p/{slug}/')
            self.assertContains(response, 'CICC — Centro Integral de Conocimiento de Calidad')
            self.assertContains(response, Resource.objects.filter(kind=slug).order_by('title').first().title)
        self.assertContains(self.client.get('/p/components/'), 'CICS → COBOL → DATACOM')

    def test_menu_groups_follow_permissions(self):
        self.login()
        participant = self.client.get('/p/dashboard/')
        self.assertContains(participant, 'Aprender')
        self.assertContains(participant, 'Conocimiento QA')
        self.assertNotContains(participant, 'Gestión y control')
        self.login('admin')
        self.assertContains(self.client.get('/p/dashboard/'), 'Gestión y control')

    def test_menu_shows_role_not_a_hardcoded_label(self):
        user = User.objects.create_user('sinrol', password=PASSWORD)
        self.client.force_login(user)
        self.assertContains(self.client.get('/p/dashboard/'), 'Sin rol')


class Dashboard(BaseCase):
    def test_uses_real_numbers_and_next_lesson(self):
        self.login()
        lessons = Lesson.objects.count()
        response = self.client.get('/p/dashboard/')
        self.assertContains(response, f'/ {lessons}')
        self.assertContains(response, 'Continúa tu ruta')
        self.assertNotContains(response, '12 lecciones')

    def test_active_attempt_is_offered_first(self):
        attempt = self.start_attempt()
        response = self.client.get('/p/dashboard/')
        self.assertContains(response, f'/attempt/{attempt.pk}/')
        self.assertContains(response, 'Reanudar evaluación')
        self.assertContains(response, '1 en curso')

    def test_recent_resources_are_remembered(self):
        self.login()
        resource = Resource.objects.filter(published=True).first()
        self.client.get(f'/resource/{resource.pk}/')
        self.assertContains(self.client.get('/p/dashboard/'), 'Vistos recientemente')


class Library(BaseCase):
    def test_search_and_protected_file(self):
        self.login()
        self.assertContains(self.client.get('/p/library/?q=batch&area=Mainframe'), 'Guía de pruebas batch')
        version = ResourceVersion.objects.exclude(file='').first()
        self.assertEqual(self.client.get(f'/download/{version.pk}/').status_code, 200)
        version.resource.published = False
        version.resource.save()
        self.assertEqual(self.client.get(f'/download/{version.pk}/').status_code, 404)

    def test_search_body_academy_and_unpublished(self):
        self.login()
        response = self.client.get('/p/search/', {'q': '980'})
        self.assertContains(response, 'JCL, COBOL, ROSCOE/TSO y Batch')
        self.assertContains(response, '/p/lessons/#lesson-')
        hidden = Resource.objects.create(title='OCULTO', description='980', area='Mainframe', published=False)
        self.assertNotContains(self.client.get('/p/search/?q=980'), hidden.title)
        self.assertEqual(self.client.get(f'/resource/{hidden.pk}/').status_code, 404)

    def test_search_highlights_and_escapes(self):
        self.login()
        self.assertContains(self.client.get('/p/search/?q=batch'), '<mark>')
        response = self.client.get('/p/search/', {'q': '<script>'})
        self.assertNotContains(response, '<script>alert')

    def test_empty_search_prompts_instead_of_listing_everything(self):
        self.login()
        self.assertContains(self.client.get('/p/search/'), 'Escribe qué quieres encontrar')

    def test_global_search_does_not_expose_answers(self):
        self.login()
        response = self.client.get('/p/search/?q=evidencia')
        self.assertNotContains(response, 'La conciliación permite explicar el destino de cada registro.')

    def test_pagination_and_filters_keep_the_query(self):
        self.login()
        for i in range(25):
            Resource.objects.create(title=f'Extra {i:02d}', description='paginacion', area='Distribuido')
        first = self.client.get('/p/library/?q=paginacion')
        self.assertContains(first, 'Página 1 de 2')
        self.assertContains(first, 'q=paginacion&amp;page=2')
        self.assertContains(self.client.get('/p/library/?q=paginacion&page=2'), 'Página 2 de 2')
        # Una página fuera de rango muestra la última, no un error.
        self.assertEqual(self.client.get('/p/library/?q=paginacion&page=99').status_code, 200)

    def test_kind_filter_and_sort(self):
        self.login()
        response = self.client.get('/p/library/?tipo=Video&orden=titulo')
        self.assertContains(response, 'Introducción a DB2 para QA')
        self.assertNotContains(response, 'Documentación de Django')

    def test_favorites(self):
        self.login()
        resource = Resource.objects.filter(published=True).first()
        self.assertContains(self.client.get('/p/favorites/'), 'Aún no tienes favoritos')
        self.client.post(f'/favorite/{resource.pk}/', {'next': '/p/library/'})
        self.assertTrue(Favorite.objects.filter(user__username='ana', resource=resource).exists())
        self.assertContains(self.client.get('/p/favorites/'), resource.title)
        self.client.post(f'/favorite/{resource.pk}/')
        self.assertFalse(Favorite.objects.exists())

    def test_favorite_redirect_rejects_external_urls(self):
        self.login()
        resource = Resource.objects.filter(published=True).first()
        response = self.client.post(f'/favorite/{resource.pk}/', {'next': 'https://evil.example/'})
        self.assertEqual(response['Location'], f'/resource/{resource.pk}/')
        self.assertEqual(self.client.get(f'/favorite/{resource.pk}/').status_code, 405)

    def test_landing_resources_open_their_course(self):
        self.login()
        resource = Resource.objects.get(title='Guía de pruebas batch y JCL')
        self.assertEqual(resource.landing, 'fundamentals')
        self.assertEqual(resource.display_title, 'Fundamentos QA')
        self.assertContains(self.client.get(f'/resource/{resource.pk}/'), 'Ruta de aprendizaje')

    def test_week_navigation(self):
        self.login()
        middle = self.client.get('/qa/fundamentos/semana/2/')
        self.assertContains(middle, 'rel="prev"')
        self.assertContains(middle, 'rel="next"')
        self.assertNotContains(self.client.get('/qa/fundamentos/semana/1/'), 'rel="prev"')
        self.assertNotContains(self.client.get('/qa/fundamentos/semana/6/'), 'rel="next"')
        self.assertEqual(self.client.get('/qa/fundamentos/semana/7/').status_code, 404)
        self.assertEqual(self.client.get('/qa/avanzado/semana/11/').status_code, 200)
        self.assertEqual(self.client.get('/qa/avanzado/semana/12/').status_code, 404)


class Learning(BaseCase):
    def test_progress_idempotent(self):
        self.login()
        lesson = Lesson.objects.first()
        for _ in range(2):
            self.client.post(f'/lesson/{lesson.pk}/complete/')
        self.assertEqual(Progress.objects.count(), 1)

    def test_completing_moves_to_next_lesson_and_celebrates_the_end(self):
        self.login()
        path = LearningPath.objects.get(title='Fundamentos QA Mainframe')
        first, second, third = path.lessons.all()
        response = self.client.post(f'/lesson/{first.pk}/complete/')
        self.assertEqual(response['Location'], f'/p/lessons/#lesson-{second.pk}')
        self.client.post(f'/lesson/{second.pk}/complete/')
        response = self.client.post(f'/lesson/{third.pk}/complete/', follow=True)
        self.assertContains(response, 'Completaste la ruta')

    def test_paths_show_progress(self):
        self.login()
        lesson = LearningPath.objects.get(title='Fundamentos QA Mainframe').lessons.first()
        self.client.post(f'/lesson/{lesson.pk}/complete/')
        self.assertContains(self.client.get('/p/paths/'), '1 de 3 lecciones')


class Exams(BaseCase):
    def test_autosave_grade_and_ownership(self):
        attempt = self.start_attempt()
        answers = {str(q.pk): q.correct for q in attempt.assessment.questions.all()}
        self.assertEqual(self.client.get(f'/attempt/{attempt.pk}/').status_code, 200)
        saved = self.client.post(f'/attempt/{attempt.pk}/save/', json.dumps(answers), content_type='application/json')
        self.assertEqual(saved.status_code, 200)
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        attempt.refresh_from_db()
        self.assertEqual(attempt.score, 100)
        self.assertIsNotNone(attempt.finished_at)
        self.assertEqual(self.client.get(f'/results/{attempt.pk}/').status_code, 200)
        self.login('diego')
        self.assertEqual(self.client.get(f'/attempt/{attempt.pk}/').status_code, 404)

    def test_expired_rejects_answers(self):
        attempt = self.start_attempt()
        attempt.deadline = timezone.now() - timedelta(seconds=1)
        attempt.save()
        response = self.client.post(f'/attempt/{attempt.pk}/save/', '{}', content_type='application/json')
        self.assertEqual(response.status_code, 409)
        attempt.refresh_from_db()
        self.assertTrue(attempt.finished)
        self.assertEqual(attempt.answers, {})

    def test_expired_attempt_is_closed_when_opened(self):
        attempt = self.start_attempt()
        attempt.deadline = timezone.now() - timedelta(seconds=5)
        attempt.save()
        response = self.client.get(f'/attempt/{attempt.pk}/')
        self.assertRedirects(response, f'/results/{attempt.pk}/')
        attempt.refresh_from_db()
        self.assertTrue(attempt.finished)

    def test_invalid_answers(self):
        attempt = self.start_attempt()
        for body in ('{"999":0}', '[1]', 'no-json', f'{{"{attempt.frozen_questions()[0]["id"]}": true}}'):
            response = self.client.post(f'/attempt/{attempt.pk}/save/', body, content_type='application/json')
            self.assertEqual(response.status_code, 400, body)

    def test_exam_page_does_not_leak_answers(self):
        attempt = self.start_attempt()
        page = self.client.get(f'/attempt/{attempt.pk}/')
        self.assertContains(page, 'role="timer"')
        self.assertContains(page, 'fieldset')
        for question in attempt.frozen_questions():
            self.assertNotContains(page, question['explanation'])

    def test_snapshot_freezes_questions_and_order_is_kept(self):
        attempt = self.start_attempt()
        original = attempt.frozen_questions()
        question = attempt.assessment.questions.first()
        answers = {str(q['id']): q['correct'] for q in original}
        question.correct = (question.correct + 1) % len(question.options)
        question.text = 'Texto cambiado después de iniciar'
        question.save()
        self.client.post(f'/attempt/{attempt.pk}/save/', json.dumps(answers), content_type='application/json')
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        attempt.refresh_from_db()
        self.assertEqual(attempt.score, 100)
        self.assertNotContains(self.client.get(f'/results/{attempt.pk}/'), 'Texto cambiado')
        self.assertEqual([q['id'] for q in attempt.frozen_questions()], [q['id'] for q in original])

    def test_only_one_active_attempt_per_person_and_assessment(self):
        first = self.start_attempt()
        assessment = first.assessment
        self.client.post(f'/exam/{assessment.pk}/start/')
        self.assertEqual(Attempt.objects.filter(user=first.user, assessment=assessment).count(), 1)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Attempt.objects.create(user=first.user, assessment=assessment, deadline=timezone.now())

    def test_attempt_limit(self):
        self.login()
        assessment = Assessment.objects.first()
        assessment.max_attempts = 1
        assessment.save()
        self.client.post(f'/exam/{assessment.pk}/start/')
        attempt = Attempt.objects.get()
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        response = self.client.post(f'/exam/{assessment.pk}/start/', follow=True)
        self.assertContains(response, 'Alcanzaste el límite de 1 intentos')
        self.assertEqual(Attempt.objects.count(), 1)
        self.assertContains(self.client.get(f'/results/{attempt.pk}/'), 'Alcanzaste el límite de intentos')

    def test_retry_and_results_summary(self):
        attempt = self.start_attempt()
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        page = self.client.get(f'/results/{attempt.pk}/')
        self.assertContains(page, 'Reintentar evaluación')
        self.assertContains(page, 'Sin responder')
        self.assertContains(page, 'Respuesta esperada')
        self.client.post(f'/exam/{attempt.assessment.pk}/start/')
        self.assertEqual(Attempt.objects.filter(user=attempt.user, finished=False).count(), 1)

    def test_history_lists_duration(self):
        attempt = self.start_attempt()
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        self.assertContains(self.client.get('/p/results/'), 'min')


class Tools(BaseCase):
    def upload(self, content):
        return self.client.post('/import/', {'file': SimpleUploadedFile('q.xlsx', content)})

    def test_excel_and_backup(self):
        self.login('lucia')
        self.assertEqual(self.client.get('/template/').status_code, 200)
        before = Question.objects.count()
        self.assertEqual(self.upload(valid_template(self.client)).status_code, 200)
        self.assertEqual(Question.objects.count(), before + 1)
        self.assertEqual(self.client.get('/report.xlsx').status_code, 200)
        self.assertEqual(self.client.post('/backup/').status_code, 200)

    def test_import_rejects_the_whole_batch_and_lists_errors(self):
        self.login('lucia')
        wb = load_workbook(BytesIO(self.client.get('/template/').content))
        wb.active.append([1, 'Rota', None, None, None, 9, 'No'])
        wb.active.append([999, 'Sin evaluación', 'a', 'b', 'c', 0, 'x'])
        wb.active.append([1, 'Opción en blanco', 'a', '   ', 'c', 0, 'x'])
        out = BytesIO()
        wb.save(out)
        before = Question.objects.count()
        response = self.upload(out.getvalue())
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Question.objects.count(), before)
        self.assertContains(response, 'Fila 3: hay celdas vacías', status_code=400)
        self.assertContains(response, 'la evaluación «999» no existe', status_code=400)
        self.assertContains(response, 'Fila 5', status_code=400)

    def test_import_rejects_non_excel_and_wrong_headers(self):
        self.login('lucia')
        self.assertContains(self.upload(b'not a zip'), 'no es un .xlsx válido', status_code=400)
        wb = load_workbook(BytesIO(self.client.get('/template/').content))
        wb.active['A1'] = 'otra_cosa'
        out = BytesIO()
        wb.save(out)
        self.assertContains(self.upload(out.getvalue()), 'Encabezados incorrectos', status_code=400)
        self.assertContains(self.client.post('/import/'), 'Selecciona un archivo', status_code=400)

    def test_import_ignores_blank_trailing_rows(self):
        self.login('lucia')
        wb = load_workbook(BytesIO(valid_template(self.client)))
        wb.active.append([None] * 7)
        out = BytesIO()
        wb.save(out)
        before = Question.objects.count()
        self.assertEqual(self.upload(out.getvalue()).status_code, 200)
        self.assertEqual(Question.objects.count(), before + 1)

    def test_report_respects_filters_and_neutralizes_formulas(self):
        evil = Assessment.objects.create(title='=HYPERLINK("http://x")', area='Mainframe')
        ana = User.objects.get(username='ana')
        diego = User.objects.get(username='diego')
        now = timezone.now()
        for user in (ana, diego):
            Attempt.objects.create(user=user, assessment=evil, deadline=now, finished=True, finished_at=now, score=80)
        self.login('lucia')
        sheet = load_workbook(BytesIO(self.client.get('/report.xlsx', {'usuario': ana.pk}).content)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], 'ana')
        self.assertTrue(rows[1][1].startswith("'="))
        everyone = load_workbook(BytesIO(self.client.get('/report.xlsx').content)).active
        self.assertEqual(len(list(everyone.iter_rows())), 3)

    def test_report_page_filters_and_summary(self):
        self.login('lucia')
        assessment = Assessment.objects.first()
        now = timezone.now()
        ana = User.objects.get(username='ana')
        Attempt.objects.create(user=ana, assessment=assessment, deadline=now, finished=True, finished_at=now, score=90)
        page = self.client.get('/p/reports/', {'usuario': ana.pk, 'desde': '2000-01-01', 'hasta': 'no-es-fecha'})
        self.assertContains(page, 'Aprobación 100%')
        self.assertEqual(self.client.get('/p/reports/', {'usuario': 'x', 'desde': '2024-13-45'}).status_code, 200)

    def test_backup_is_restorable_and_has_no_passwords(self):
        Favorite.objects.create(user=User.objects.get(username='ana'), resource=Resource.objects.first())
        now = timezone.now()
        Attempt.objects.create(
            user=User.objects.get(username='ana'), assessment=Assessment.objects.first(), deadline=now, finished=True
        )
        self.login('admin')
        content = self.client.post('/backup/').content.decode()
        data = json.loads(content)
        models = {item['model'] for item in data}
        self.assertTrue({'auth.user', 'auth.group', 'hub.resource', 'hub.attempt', 'hub.favorite'} <= models)
        for item in data:
            if item['model'] == 'auth.user':
                self.assertNotIn('password', item['fields'])
        restored = list(serializers.deserialize('json', content))
        self.assertEqual(len(restored), len(data))

    def test_seed_repeat_and_transversal_validation(self):
        counts = (Resource.objects.count(), Lesson.objects.count())
        call_command('seed_demo', password=PASSWORD)
        self.assertEqual(counts, (Resource.objects.count(), Lesson.objects.count()))
        for resource in Resource.objects.filter(area='Transversal'):
            resource.full_clean()
        self.assertEqual(len(CONTENT), 26)


class Login(BaseCase):
    def attempt(self, password):
        return self.client.post('/login/', {'username': 'ana', 'password': password})

    def test_too_many_failures_lock_the_account_temporarily(self):
        for _ in range(5):
            self.assertContains(self.attempt('mala'), 'Usuario o contraseña incorrectos')
        locked = self.attempt(PASSWORD)
        self.assertEqual(locked.status_code, 429)
        self.assertContains(locked, 'Demasiados intentos', status_code=429)

    def test_success_resets_the_counter(self):
        for _ in range(3):
            self.attempt('mala')
        self.assertEqual(self.attempt(PASSWORD).status_code, 302)
        for _ in range(4):
            self.attempt('mala')
        self.assertEqual(self.attempt(PASSWORD).status_code, 302)

    @override_settings(SHOW_DEMO_CREDENTIALS=False)
    def test_demo_credentials_can_be_hidden(self):
        self.assertNotContains(self.client.get('/login/'), 'demo-credentials.txt')

    @override_settings(SHOW_DEMO_CREDENTIALS=True)
    def test_demo_credentials_shown_when_enabled(self):
        self.assertContains(self.client.get('/login/'), 'demo-credentials.txt')

    def test_logout_requires_post(self):
        self.login()
        self.assertEqual(self.client.get('/logout/').status_code, 405)
        self.assertEqual(self.client.post('/logout/').status_code, 302)


class Settings(TestCase):
    def run_settings(self, **env):
        base = {k: v for k, v in os.environ.items() if k not in ('DEBUG', 'DJANGO_SECRET_KEY', 'VERCEL')}
        base.update({'DJANGO_SETTINGS_MODULE': 'config.settings', 'QA_SQLITE': '1', **env})
        return subprocess.run(
            [
                sys.executable,
                '-c',
                'import django; django.setup(); from django.conf import settings; print(settings.DEBUG)',
            ],
            cwd=settings.BASE_DIR,
            env=base,
            capture_output=True,
            text=True,
        )

    def test_production_requires_a_real_secret_key(self):
        result = self.run_settings(DEBUG='0')
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('DJANGO_SECRET_KEY', result.stderr)

    def test_production_with_secret_key_starts(self):
        result = self.run_settings(DEBUG='0', DJANGO_SECRET_KEY='x' * 50)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip(), 'False')

    def test_vercel_is_not_debug_by_default(self):
        self.assertNotEqual(self.run_settings(VERCEL='1').returncode, 0)


class DemoAccounts(BaseCase):
    def test_random_passwords_are_generated_and_saved_to_a_file(self):
        folder = Path(tempfile.mkdtemp(prefix='cicc-creds-'))
        with override_settings(BASE_DIR=folder):
            call_command('seed_demo', '--reset-passwords')
        lines = (folder / 'demo-credentials.txt').read_text(encoding='utf-8').splitlines()
        creds = dict(line.split(': ', 1) for line in lines if ': ' in line)
        self.assertEqual(set(creds), {'admin', 'ana', 'diego', 'lucia'})
        self.assertNotIn(PASSWORD, creds.values())
        self.assertEqual(len(set(creds.values())), 4)
        self.assertTrue(self.client.login(username='ana', password=creds['ana']))
        self.assertFalse(self.client.login(username='ana', password=PASSWORD))

    def test_existing_passwords_are_kept_without_reset(self):
        call_command('seed_demo')
        self.assertTrue(self.client.login(username='ana', password=PASSWORD))


class QuestionEditor(BaseCase):
    def payload(self, **overrides):
        data = {
            'assessment': Assessment.objects.first().pk,
            'kind': 'single',
            'topic': 'Tema nuevo',
            'text': '¿Pregunta del editor?',
            'explanation': 'Porque sí.',
            'options_text': 'Uno\nDos\nTres',
            'correct_number': 2,
        }
        data.update(overrides)
        return data

    def test_participants_cannot_use_the_editor(self):
        self.login()
        self.assertEqual(self.client.get('/question/new/').status_code, 403)
        self.assertEqual(self.client.post('/question/1/delete/').status_code, 403)

    def test_create_edit_and_delete(self):
        self.login('lucia')
        self.assertContains(self.client.get('/question/new/'), 'Vista previa')
        response = self.client.post('/question/new/', self.payload())
        self.assertRedirects(response, '/p/questions/')
        question = Question.objects.get(text='¿Pregunta del editor?')
        self.assertEqual(
            (question.options, question.correct, question.topic), (['Uno', 'Dos', 'Tres'], 1, 'Tema nuevo')
        )
        self.assertContains(self.client.get(f'/question/{question.pk}/edit/'), 'Uno')
        self.client.post(f'/question/{question.pk}/edit/', self.payload(options_text='A\nB', correct_number=1))
        question.refresh_from_db()
        self.assertEqual((question.options, question.correct), (['A', 'B'], 0))
        self.client.post(f'/question/{question.pk}/delete/')
        self.assertFalse(Question.objects.filter(pk=question.pk).exists())

    def test_validation_errors_are_shown_and_nothing_is_saved(self):
        self.login('lucia')
        before = Question.objects.count()
        for bad in (
            {'options_text': 'Solo una'},
            {'options_text': 'a\na'},
            {'options_text': 'a\nb', 'correct_number': 3},
            {'options_text': 'a\nb\nc\nd\ne\nf\ng'},
            {'text': ''},
        ):
            response = self.client.post('/question/new/', self.payload(**bad))
            self.assertEqual(response.status_code, 200, bad)
            self.assertContains(response, 'Revisa los campos marcados')
        self.assertEqual(Question.objects.count(), before)

    def test_editing_does_not_change_running_attempts(self):
        attempt = self.start_attempt()
        question = attempt.assessment.questions.first()
        self.login('lucia')
        self.client.post(f'/question/{question.pk}/delete/')
        attempt.refresh_from_db()
        self.assertIn(question.pk, [q['id'] for q in attempt.frozen_questions()])

    def test_list_filter_and_actions(self):
        self.login('lucia')
        page = self.client.get('/p/questions/', {'evaluacion': Assessment.objects.first().pk})
        self.assertContains(page, 'Nueva pregunta')
        self.assertContains(page, '/edit/')
        self.assertContains(page, 'Evidencia y conciliación')


class Topics(BaseCase):
    def test_results_break_down_by_topic(self):
        attempt = self.start_attempt()
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        page = self.client.get(f'/results/{attempt.pk}/')
        self.assertContains(page, 'Desempeño por tema')
        self.assertContains(page, 'Pruebas negativas')
        self.assertContains(page, 'Reforzar')

    def test_excel_topic_column_is_optional_and_saved(self):
        self.login('lucia')
        template = valid_template(self.client)
        self.client.post('/import/', {'file': SimpleUploadedFile('q.xlsx', template)})
        self.assertEqual(Question.objects.filter(topic='Pruebas smoke').count(), 1)
        wb = load_workbook(BytesIO(template))
        wb.active.delete_cols(8)
        out = BytesIO()
        wb.save(out)
        self.assertEqual(
            self.client.post('/import/', {'file': SimpleUploadedFile('s.xlsx', out.getvalue())}).status_code, 200
        )


class Certificates(BaseCase):
    def finish(self, correct=True):
        attempt = self.start_attempt()
        answers = {
            str(q['id']): (q['correct'] if correct else (q['correct'] + 1) % len(q['options']))
            for q in attempt.frozen_questions()
        }
        self.client.post(f'/attempt/{attempt.pk}/save/', json.dumps(answers), content_type='application/json')
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        attempt.refresh_from_db()
        return attempt

    def test_passing_issues_one_certificate_with_a_public_page(self):
        attempt = self.finish()
        certificate = attempt.certificate
        self.assertRegex(certificate.code, r'^CICC-[0-9A-F]{10}$')
        self.client.logout()
        page = self.client.get(f'/certificate/{certificate.code}/')
        self.assertContains(page, 'Certificado de aprobación')
        self.assertContains(page, 'Ana Torres')
        self.assertContains(page, attempt.assessment.title)
        self.assertNotContains(page, 'Mis certificados')

    def test_failing_does_not_issue_one(self):
        attempt = self.finish(correct=False)
        self.assertFalse(Certificate.objects.filter(attempt=attempt).exists())
        self.assertNotContains(self.client.get(f'/results/{attempt.pk}/'), 'Obtuviste un certificado')

    def test_unknown_code_is_404_and_listing_and_result_link(self):
        self.assertEqual(self.client.get('/certificate/CICC-0000000000/').status_code, 404)
        attempt = self.finish()
        self.login()
        self.assertContains(self.client.get(f'/results/{attempt.pk}/'), 'Obtuviste un certificado')
        self.assertContains(self.client.get('/p/certificates/'), attempt.certificate.code)
        self.assertContains(self.client.get(f'/certificate/{attempt.certificate.code}/'), 'Mis certificados')

    def test_empty_certificates_page(self):
        self.login()
        self.assertContains(self.client.get('/p/certificates/'), 'Todavía no tienes certificados')


class Cohorts(BaseCase):
    def test_assigned_groups_limit_who_sees_and_starts(self):
        assessment = Assessment.objects.first()
        assessment.groups.add(Group.objects.get(name='Evaluador'))
        self.login()
        self.assertNotContains(self.client.get('/p/assessments/'), assessment.title)
        self.assertEqual(self.client.post(f'/exam/{assessment.pk}/start/').status_code, 404)
        self.login('lucia')
        self.assertContains(self.client.get('/p/assessments/'), assessment.title)
        self.login('admin')
        self.assertContains(self.client.get('/p/assessments/'), assessment.title)

    def test_open_and_close_dates(self):
        assessment = Assessment.objects.first()
        self.login()
        assessment.opens_at = timezone.now() + timedelta(days=1)
        assessment.save()
        self.assertContains(self.client.get('/p/assessments/'), 'Disponible desde el')
        response = self.client.post(f'/exam/{assessment.pk}/start/', follow=True)
        self.assertContains(response, 'no está disponible')
        self.assertFalse(Attempt.objects.exists())
        assessment.opens_at = None
        assessment.closes_at = timezone.now() - timedelta(hours=1)
        assessment.save()
        self.assertContains(self.client.get('/p/assessments/'), 'Cerrada desde el')

    def test_a_running_attempt_can_continue_after_closing_date(self):
        attempt = self.start_attempt()
        attempt.assessment.closes_at = timezone.now() - timedelta(minutes=1)
        attempt.assessment.save()
        self.assertEqual(self.client.get(f'/attempt/{attempt.pk}/').status_code, 200)


class EmailAndExpiry(BaseCase):
    def test_close_expired_command_grades_and_emails(self):
        User.objects.filter(username='ana').update(email='ana@example.com')
        attempt = self.start_attempt()
        attempt.deadline = timezone.now() - timedelta(minutes=1)
        attempt.save()
        with self.captureOnCommitCallbacks(execute=True):
            call_command('close_expired')
        attempt.refresh_from_db()
        self.assertTrue(attempt.finished)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ['ana@example.com'])
        self.assertIn('Resultado de', mail.outbox[0].subject)

    def test_no_email_without_an_address(self):
        attempt = self.start_attempt()
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(f'/attempt/{attempt.pk}/finish/')
        self.assertEqual(len(mail.outbox), 0)

    def test_finishing_sends_the_result_with_certificate_link(self):
        User.objects.filter(username='ana').update(email='ana@example.com')
        attempt = self.start_attempt()
        answers = {str(q['id']): q['correct'] for q in attempt.frozen_questions()}
        self.client.post(f'/attempt/{attempt.pk}/save/', json.dumps(answers), content_type='application/json')
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(f'/attempt/{attempt.pk}/finish/')
        self.assertIn('/certificate/CICC-', mail.outbox[0].body)


class PasswordFlows(BaseCase):
    def test_reset_by_email_end_to_end(self):
        User.objects.filter(username='ana').update(email='ana@example.com')
        self.assertContains(self.client.get('/login/'), 'Olvidaste tu contraseña')
        response = self.client.post('/password/reset/', {'email': 'ana@example.com'})
        self.assertRedirects(response, '/password/reset/enviado/')
        self.assertEqual(len(mail.outbox), 1)
        link = re.search(r'https?://[^/]+(/password/reset/[^\s]+)', mail.outbox[0].body).group(1)
        confirm = self.client.get(link, follow=True)
        self.assertContains(confirm, 'Crea tu nueva contraseña')
        done = self.client.post(
            confirm.request['PATH_INFO'], {'new_password1': 'NuevaClave-2026!x', 'new_password2': 'NuevaClave-2026!x'}
        )
        self.assertRedirects(done, '/password/reset/completo/')
        self.assertTrue(self.client.login(username='ana', password='NuevaClave-2026!x'))

    def test_reset_does_not_reveal_whether_the_email_exists(self):
        response = self.client.post('/password/reset/', {'email': 'nadie@example.com'})
        self.assertRedirects(response, '/password/reset/enviado/')
        self.assertEqual(len(mail.outbox), 0)

    def test_bad_reset_link(self):
        self.assertContains(self.client.get('/password/reset/abc/xyz/'), 'Este enlace ya no sirve')

    def test_change_password_inside_the_app(self):
        self.assertEqual(self.client.get('/cuenta/password/').status_code, 302)
        self.login()
        self.assertContains(self.client.get('/cuenta/password/'), 'Cambiar contraseña')
        bad = self.client.post(
            '/cuenta/password/',
            {'old_password': 'mala', 'new_password1': 'OtraClave-2026!x', 'new_password2': 'OtraClave-2026!x'},
        )
        self.assertEqual(bad.status_code, 200)
        self.assertTrue(self.client.login(username='ana', password=PASSWORD))
        ok = self.client.post(
            '/cuenta/password/',
            {'old_password': PASSWORD, 'new_password1': 'OtraClave-2026!x', 'new_password2': 'OtraClave-2026!x'},
        )
        self.assertRedirects(ok, '/cuenta/password/')
        self.assertTrue(self.client.login(username='ana', password='OtraClave-2026!x'))
        self.assertTrue(Audit.objects.filter(action='Cambio de contraseña').exists())

    def test_weak_passwords_are_rejected(self):
        self.login()
        weak = self.client.post(
            '/cuenta/password/', {'old_password': PASSWORD, 'new_password1': '12345678', 'new_password2': '12345678'}
        )
        self.assertEqual(weak.status_code, 200)


class CourseWeeks(BaseCase):
    def test_week_completion_and_progress(self):
        self.login()
        self.assertContains(self.client.get('/qa/fundamentos/'), '0 de 6 semanas completadas')
        response = self.client.post('/curso/fundamentals/semana/2/completar/', follow=True)
        self.assertContains(response, 'Semana completada')
        self.assertContains(self.client.get('/qa/fundamentos/'), '1 de 6 semanas completadas')
        self.assertContains(self.client.get('/p/dashboard/'), '1 de 6 semanas')
        self.client.post('/curso/fundamentals/semana/2/completar/')
        self.assertFalse(WeekProgress.objects.exists())

    def test_invalid_week_or_course(self):
        self.login()
        self.assertEqual(self.client.post('/curso/fundamentals/semana/9/completar/').status_code, 404)
        self.assertEqual(self.client.post('/curso/otro/semana/1/completar/').status_code, 404)
        self.assertEqual(self.client.get('/curso/fundamentals/semana/1/completar/').status_code, 405)

    def test_database_content_overrides_the_file_and_can_be_exported(self):
        self.login()
        original = self.client.get('/qa/fundamentos/semana/1/')
        self.assertContains(original, 'QA asegura la calidad')
        CourseWeek.objects.create(
            course='fundamentals', number=1, title='Título editado', body='<h2>Contenido editado</h2>'
        )
        edited = self.client.get('/qa/fundamentos/semana/1/')
        self.assertContains(edited, 'Contenido editado')
        self.assertContains(edited, 'Título editado')
        self.assertNotContains(edited, 'QA asegura la calidad')
        call_command('export_weeks')
        self.assertEqual(CourseWeek.objects.count(), 17)
        self.assertEqual(CourseWeek.objects.get(course='fundamentals', number=1).body, '<h2>Contenido editado</h2>')
        call_command('export_weeks', '--overwrite')
        self.assertIn('QA asegura la calidad', CourseWeek.objects.get(course='fundamentals', number=1).body)


class VideoPlayer(BaseCase):
    CONTENT = bytes(range(256)) * 8

    def make_version(self, name='demo.mp4'):
        resource = Resource.objects.create(title='Video propio', description='d', area='Mainframe', kind='Video')
        return ResourceVersion.objects.create(
            resource=resource, number=1, notes='n', file=SimpleUploadedFile(name, self.CONTENT)
        )

    def test_full_and_ranged_responses(self):
        version = self.make_version()
        self.login()
        full = self.client.get(f'/stream/{version.pk}/')
        self.assertEqual(full.status_code, 200)
        self.assertEqual(full['Accept-Ranges'], 'bytes')
        self.assertEqual(b''.join(full.streaming_content), self.CONTENT)
        part = self.client.get(f'/stream/{version.pk}/', HTTP_RANGE='bytes=10-19')
        self.assertEqual(part.status_code, 206)
        self.assertEqual(part['Content-Range'], f'bytes 10-19/{len(self.CONTENT)}')
        self.assertEqual(b''.join(part.streaming_content), self.CONTENT[10:20])
        tail = self.client.get(f'/stream/{version.pk}/', HTTP_RANGE='bytes=-5')
        self.assertEqual(b''.join(tail.streaming_content), self.CONTENT[-5:])
        open_ended = self.client.get(f'/stream/{version.pk}/', HTTP_RANGE=f'bytes={len(self.CONTENT) - 3}-')
        self.assertEqual(b''.join(open_ended.streaming_content), self.CONTENT[-3:])
        self.assertEqual(self.client.get(f'/stream/{version.pk}/', HTTP_RANGE='bytes=99999-').status_code, 416)

    def test_access_rules(self):
        version = self.make_version()
        self.assertEqual(self.client.get(f'/stream/{version.pk}/').status_code, 302)
        self.login()
        document = ResourceVersion.objects.exclude(file='').first()
        self.assertEqual(self.client.get(f'/stream/{document.pk}/').status_code, 404)
        version.resource.published = False
        version.resource.save()
        self.assertEqual(self.client.get(f'/stream/{version.pk}/').status_code, 404)

    def test_resource_page_embeds_the_player(self):
        version = self.make_version()
        self.login()
        page = self.client.get(f'/resource/{version.resource.pk}/')
        self.assertContains(page, '<video')
        self.assertContains(page, f'/stream/{version.pk}/')

    def test_upload_validators(self):
        version = self.make_version()
        version.file = SimpleUploadedFile('virus.exe', b'x')
        with self.assertRaises(ValidationError):
            version.full_clean()
        version.file = SimpleUploadedFile('ok.pdf', b'x')
        version.full_clean()


class Appearance(BaseCase):
    def test_theme_toggle_is_present_everywhere(self):
        self.login()
        page = self.client.get('/p/dashboard/')
        self.assertContains(page, 'id="theme-toggle"')
        self.assertContains(page, 'theme.js')
        self.assertContains(self.client.get('/login/'), 'theme.js')

    def test_menu_has_certificates_and_password_link(self):
        self.login()
        page = self.client.get('/p/dashboard/')
        self.assertContains(page, 'Mis certificados')
        self.assertContains(page, 'href="/cuenta/"')


def practice_assessment():
    return Assessment.objects.get(title='Práctica: tipos de pregunta')


def perfect_answers(attempt):
    answers = {}
    for q in attempt.frozen_questions():
        kind = q.get('kind', 'single')
        answers[str(q['id'])] = (
            q['correct_set'] if kind == 'multiple' else q['accepted'][0] if kind == 'short' else q['correct']
        )
    return answers


class QuestionKinds(BaseCase):
    def save(self, attempt, answers):
        return self.client.post(f'/attempt/{attempt.pk}/save/', json.dumps(answers), content_type='application/json')

    def test_model_validation_per_kind(self):
        base = {'assessment': practice_assessment(), 'text': 't', 'explanation': 'e'}
        Question(kind='truefalse', options=['Verdadero', 'Falso'], correct=1, **base).full_clean()
        Question(kind='multiple', options=['a', 'b', 'c'], correct_set=[0, 2], **base).full_clean()
        Question(kind='short', options=[], accepted=['x'], **base).full_clean()
        for bad in (
            Question(kind='truefalse', options=['Sí', 'No'], correct=0, **base),
            Question(kind='multiple', options=['a', 'b'], correct_set=[], **base),
            Question(kind='multiple', options=['a', 'b'], correct_set=[0, 0], **base),
            Question(kind='multiple', options=['a', 'b'], correct_set=[5], **base),
            Question(kind='short', options=[], accepted=[], **base),
            Question(kind='short', options=[], accepted=['  '], **base),
            Question(kind='single', options=['a', 'b'], correct=4, **base),
        ):
            with self.assertRaises(ValidationError):
                bad.full_clean()

    def test_all_kinds_graded_correctly(self):
        attempt = self.start_attempt(practice_assessment())
        self.save(attempt, perfect_answers(attempt))
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        attempt.refresh_from_db()
        self.assertEqual(attempt.score, 100)
        self.assertTrue(hasattr(attempt, 'certificate'))

    def test_multiple_needs_exactly_the_right_set(self):
        attempt = self.start_attempt(practice_assessment())
        multiple = next(q for q in attempt.frozen_questions() if q['kind'] == 'multiple')
        for wrong in ([0, 1], [0, 1, 2, 3], [], [3]):
            attempt.answers = {str(multiple['id']): wrong}
            self.assertFalse(is_correct(multiple, attempt.answers[str(multiple['id'])]), wrong)
        self.assertTrue(is_correct(multiple, list(reversed(multiple['correct_set']))))

    def test_short_answer_ignores_case_accents_and_spacing(self):
        short = {'kind': 'short', 'accepted': ['regresión']}
        for given in ('Regresion', '  REGRESIÓN ', 'regresión'):
            self.assertTrue(is_correct(short, given), given)
        for given in ('regresiones', '', None, 5):
            self.assertFalse(is_correct(short, given), given)

    def test_answer_shapes_are_validated(self):
        attempt = self.start_attempt(practice_assessment())
        by_kind = {q['kind']: q for q in attempt.frozen_questions()}
        ids = {k: str(v['id']) for k, v in by_kind.items()}
        good = {ids['multiple']: [0, 1], ids['short']: 'regresion', ids['truefalse']: 1, ids['single']: 0}
        self.assertEqual(self.save(attempt, good).status_code, 200)
        for bad in (
            {ids['multiple']: 1},
            {ids['multiple']: [0, 0]},
            {ids['multiple']: [9]},
            {ids['multiple']: ['0']},
            {ids['short']: 5},
            {ids['short']: 'x' * 201},
            {ids['single']: [0]},
            {ids['single']: 7},
            {ids['truefalse']: 2},
        ):
            self.assertEqual(self.save(attempt, bad).status_code, 400, bad)

    def test_exam_markup_per_kind_without_leaking_answers(self):
        attempt = self.start_attempt(practice_assessment())
        page = self.client.get(f'/attempt/{attempt.pk}/')
        self.assertContains(page, 'type="checkbox"')
        self.assertContains(page, 'type="text"')
        self.assertContains(page, 'data-kind="multiple"')
        self.assertContains(page, 'RESPUESTA CORTA')
        self.assertNotContains(page, 'pruebas de regresión')
        self.assertNotContains(page, 'correct_set')

    def test_results_show_each_kind(self):
        attempt = self.start_attempt(practice_assessment())
        by_kind = {q['kind']: q for q in attempt.frozen_questions()}
        self.save(attempt, {str(by_kind['short']['id']): 'otra cosa', str(by_kind['multiple']['id']): [0]})
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        page = self.client.get(f'/results/{attempt.pk}/')
        self.assertContains(page, 'Respuestas aceptadas')
        self.assertContains(page, 'otra cosa')
        self.assertContains(page, 'Selección múltiple: había que marcar exactamente')

    def test_editor_creates_each_kind(self):
        self.login('lucia')
        base = {'assessment': practice_assessment().pk, 'topic': 'X', 'explanation': 'e'}
        cases = [
            ({'kind': 'multiple', 'text': 'M?', 'options_text': 'a\nb\nc', 'correct_numbers_text': '1, 3'}, 'M?'),
            ({'kind': 'truefalse', 'text': 'V?', 'correct_number': 2}, 'V?'),
            ({'kind': 'short', 'text': 'S?', 'accepted_text': 'uno\ndos'}, 'S?'),
        ]
        for extra, _text in cases:
            response = self.client.post('/question/new/', {**base, **extra})
            self.assertRedirects(response, '/p/questions/')
        multiple = Question.objects.get(text='M?')
        self.assertEqual((multiple.options, multiple.correct_set), (['a', 'b', 'c'], [0, 2]))
        truefalse = Question.objects.get(text='V?')
        self.assertEqual((truefalse.options, truefalse.correct), (['Verdadero', 'Falso'], 1))
        short = Question.objects.get(text='S?')
        self.assertEqual((short.options, short.accepted), ([], ['uno', 'dos']))
        self.assertContains(self.client.get(f'/question/{multiple.pk}/edit/'), '1, 3')

    def test_editor_rejects_bad_input_per_kind(self):
        self.login('lucia')
        base = {'assessment': practice_assessment().pk, 'topic': '', 'explanation': 'e', 'text': 'T'}
        before = Question.objects.count()
        for bad in (
            {'kind': 'multiple', 'options_text': 'a\nb', 'correct_numbers_text': ''},
            {'kind': 'multiple', 'options_text': 'a\nb', 'correct_numbers_text': '1, 1'},
            {'kind': 'multiple', 'options_text': 'a\nb', 'correct_numbers_text': '3'},
            {'kind': 'multiple', 'options_text': 'a\nb', 'correct_numbers_text': 'x'},
            {'kind': 'truefalse', 'correct_number': 3},
            {'kind': 'truefalse'},
            {'kind': 'short', 'accepted_text': '  \n '},
        ):
            response = self.client.post('/question/new/', {**base, **bad})
            self.assertEqual(response.status_code, 200, bad)
            self.assertContains(response, 'Revisa los campos marcados')
        self.assertEqual(Question.objects.count(), before)


class RandomDraw(BaseCase):
    def test_draw_balances_topics_and_respects_count(self):
        questions = list(practice_assessment().questions.all())
        self.assertEqual(len(draw_questions(questions, 0)), 4)
        self.assertEqual(len(draw_questions(questions, 99)), 4)
        topics = {q.topic for q in questions}
        picked = draw_questions(questions, 3)
        self.assertEqual(len(picked), 3)
        self.assertEqual(len({q.topic for q in picked}), 3)
        self.assertTrue({q.topic for q in picked} <= topics)

    def test_attempts_receive_only_the_configured_number(self):
        assessment = practice_assessment()
        assessment.questions_per_attempt = 2
        assessment.save()
        attempt = self.start_attempt(assessment)
        self.assertEqual(len(attempt.frozen_questions()), 2)
        self.assertContains(self.client.get(f'/attempt/{attempt.pk}/'), 'DE 2')
        # Cada intento puede recibir preguntas distintas; la nota solo usa las suyas.
        self.save_all_and_finish(attempt)
        attempt.refresh_from_db()
        self.assertEqual(attempt.score, 100)

    def save_all_and_finish(self, attempt):
        self.client.post(
            f'/attempt/{attempt.pk}/save/', json.dumps(perfect_answers(attempt)), content_type='application/json'
        )
        self.client.post(f'/attempt/{attempt.pk}/finish/')


class LessonChecks(BaseCase):
    def lesson(self):
        return Lesson.objects.get(path__title='Academia CICC: certificación integral', position=2)

    def test_lesson_with_checks_needs_correct_answers(self):
        self.login()
        lesson = self.lesson()
        check = lesson.checks.get()
        self.assertContains(self.client.get('/p/lessons/'), 'Comprueba lo aprendido')
        for data in ({}, {f'check-{check.pk}': '2'}, {f'check-{check.pk}': 'x'}):
            response = self.client.post(f'/lesson/{lesson.pk}/complete/', data, follow=True)
            self.assertContains(response, 'Todavía no')
            self.assertFalse(Progress.objects.filter(lesson=lesson).exists())
        self.assertContains(response, check.explanation)
        self.client.post(f'/lesson/{lesson.pk}/complete/', {f'check-{check.pk}': '0'})
        self.assertTrue(Progress.objects.filter(user__username='ana', lesson=lesson).exists())

    def test_lessons_without_checks_keep_working(self):
        self.login()
        lesson = Lesson.objects.filter(checks__isnull=True).first()
        self.client.post(f'/lesson/{lesson.pk}/complete/')
        self.assertTrue(Progress.objects.filter(lesson=lesson).exists())

    def test_answers_are_not_in_the_page_source(self):
        self.login()
        check = self.lesson().checks.get()
        page = self.client.get('/p/lessons/')
        self.assertNotContains(page, check.explanation)


class TagsAndRelated(BaseCase):
    def test_tag_filter_is_exact_not_substring(self):
        self.login()
        Resource.objects.create(title='Rápido', description='d', area='Mainframe', tags='rapid')
        page = self.client.get('/p/library/?etiqueta=api')
        self.assertContains(page, 'Checklist de pruebas de API REST')
        self.assertContains(page, 'Etiqueta:')
        self.assertNotContains(page, 'Rápido')
        self.assertNotContains(page, 'Documentación de Django')
        self.assertContains(self.client.get('/p/library/?etiqueta=PYTHON'), 'Documentación de Django')

    def test_resource_page_shows_tags_and_related(self):
        self.login()
        batch = Resource.objects.get(title='Guía de pruebas batch y JCL')
        self.assertEqual(batch.tag_list, ['jcl', 'batch', 'mainframe'])
        db2 = Resource.objects.get(title='Introducción a DB2 para QA')
        page = self.client.get(f'/resource/{db2.pk}/')
        self.assertContains(page, 'Ver también')
        self.assertContains(page, '/p/library/?etiqueta=db2')
        # La relación explícita es simétrica.
        self.assertIn(batch, db2.related.all())

    def test_same_tag_suggestions_when_nothing_is_related(self):
        self.login()
        django = Resource.objects.get(title='Documentación de Django')
        self.assertContains(self.client.get(f'/resource/{django.pk}/'), 'Automatización de servicios con Python')

    def test_unpublished_related_resources_are_hidden(self):
        self.login()
        db2 = Resource.objects.get(title='Introducción a DB2 para QA')
        batch = Resource.objects.get(title='Guía de pruebas batch y JCL')
        batch.published = False
        batch.save()
        self.assertNotContains(self.client.get(f'/resource/{db2.pk}/'), 'Fundamentos QA')


class Uploads(BaseCase):
    def resource(self):
        return Resource.objects.get(title='Documentación de Django')

    def test_editor_uploads_a_new_version(self):
        self.login('diego')
        resource = self.resource()
        self.assertContains(self.client.get(f'/resource/{resource.pk}/'), 'Subir una versión nueva')
        file = SimpleUploadedFile('manual.pdf', b'%PDF-1.4 demo')
        response = self.client.post(
            f'/resource/{resource.pk}/upload/', {'file': file, 'notes': 'Nueva guía'}, follow=True
        )
        self.assertContains(response, 'Se publicó la versión 3')
        version = resource.versions.get(number=3)
        self.assertEqual(version.notes, 'Nueva guía')
        self.assertTrue(Audit.objects.filter(action__startswith='Sube versión 3').exists())

    def test_rejects_wrong_type_missing_file_and_participants(self):
        resource = self.resource()
        self.login('diego')
        bad = self.client.post(
            f'/resource/{resource.pk}/upload/', {'file': SimpleUploadedFile('x.exe', b'MZ'), 'notes': 'n'}, follow=True
        )
        self.assertContains(bad, 'No se subió el archivo')
        missing = self.client.post(f'/resource/{resource.pk}/upload/', {'notes': 'n'}, follow=True)
        self.assertContains(missing, 'No se subió el archivo')
        self.assertEqual(resource.versions.count(), 2)
        self.login()
        page = self.client.get(f'/resource/{resource.pk}/')
        self.assertNotContains(page, 'Subir una versión nueva')
        self.assertEqual(self.client.post(f'/resource/{resource.pk}/upload/', {'notes': 'n'}).status_code, 403)
        self.assertEqual(self.client.get(f'/resource/{resource.pk}/upload/').status_code, 403)


class Profile(BaseCase):
    def test_profile_page_and_update(self):
        self.login()
        page = self.client.get('/cuenta/')
        self.assertContains(page, 'Mi perfil')
        self.assertContains(page, 'Progreso por ruta')
        response = self.client.post(
            '/cuenta/', {'first_name': 'Ana María', 'last_name': 'Torres', 'email': 'ana@example.com'}
        )
        self.assertRedirects(response, '/cuenta/')
        ana = User.objects.get(username='ana')
        self.assertEqual((ana.first_name, ana.email), ('Ana María', 'ana@example.com'))
        self.assertTrue(Audit.objects.filter(action='Actualiza su perfil').exists())

    def test_invalid_email_and_anonymous(self):
        self.assertEqual(self.client.get('/cuenta/').status_code, 302)
        self.login()
        bad = self.client.post('/cuenta/', {'first_name': 'A', 'last_name': 'B', 'email': 'no-es-correo'})
        self.assertContains(bad, 'Revisa los campos marcados')
        self.assertEqual(User.objects.get(username='ana').email, '')

    def test_profile_shows_history_scores_and_certificates(self):
        attempt = self.start_attempt()
        answers = {str(q['id']): q['correct'] for q in attempt.frozen_questions()}
        self.client.post(f'/attempt/{attempt.pk}/save/', json.dumps(answers), content_type='application/json')
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        page = self.client.get('/cuenta/')
        self.assertContains(page, 'Mejores puntajes')
        self.assertContains(page, attempt.assessment.title)
        self.assertContains(page, 'Ver mis certificados')

    def test_sidebar_links_to_profile(self):
        self.login()
        self.assertContains(self.client.get('/p/dashboard/'), 'href="/cuenta/"')


class ReportsCharts(BaseCase):
    def finish(self, correct=True):
        attempt = self.start_attempt()
        answers = {
            str(q['id']): (q['correct'] if correct else (q['correct'] + 1) % len(q['options']))
            for q in attempt.frozen_questions()
        }
        self.client.post(f'/attempt/{attempt.pk}/save/', json.dumps(answers), content_type='application/json')
        self.client.post(f'/attempt/{attempt.pk}/finish/')
        return attempt

    def test_charts_tables_and_hardest_questions(self):
        self.finish(correct=False)
        self.login('lucia')
        page = self.client.get('/p/reports/')
        self.assertContains(page, 'id="passChart"')
        self.assertContains(page, 'id="trendChart"')
        self.assertContains(page, 'chart-reports')
        self.assertContains(page, 'Preguntas que más se fallan')
        self.assertContains(page, '0%')
        data = page.context['charts']
        self.assertEqual(data['pass']['values'], [0])
        self.assertEqual(data['trend']['attempts'], [1])
        self.assertEqual(page.context['hardest'][0]['pct'], 0)

    def test_hardest_uses_the_questions_each_person_saw(self):
        attempt = self.finish(correct=True)
        Question.objects.filter(assessment=attempt.assessment).update(text='Cambiada después')
        self.login('lucia')
        page = self.client.get('/p/reports/')
        texts = [row['text'] for row in page.context['hardest']]
        self.assertTrue(texts)
        self.assertNotIn('Cambiada después', texts)

    def test_empty_report_still_renders(self):
        self.login('lucia')
        page = self.client.get('/p/reports/')
        self.assertContains(page, 'Sin intentos con estos filtros')
        self.assertEqual(page.context['charts']['pass']['labels'], [])


class AuditFilters(BaseCase):
    def test_filters_and_csv(self):
        User.objects.get(username='ana')
        Audit.objects.create(user=User.objects.get(username='ana'), action='Descarga recurso 1')
        Audit.objects.create(user=User.objects.get(username='diego'), action='=HYPERLINK("http://x")')
        self.login('lucia')
        page = self.client.get('/p/audit/', {'q': 'descarga'})
        self.assertContains(page, '<mark>Descarga</mark> recurso 1')
        self.assertNotContains(page, 'HYPERLINK')
        ana = User.objects.get(username='ana')
        only_ana = self.client.get('/p/audit/', {'usuario': ana.pk})
        self.assertNotContains(only_ana, 'HYPERLINK')
        future = self.client.get('/p/audit/', {'desde': '2999-01-01'})
        self.assertContains(future, 'No hay eventos con estos filtros')
        self.assertEqual(self.client.get('/p/audit/', {'desde': '2024-13-99'}).status_code, 200)

    def test_csv_export_respects_filters_and_neutralizes_formulas(self):
        Audit.objects.create(user=User.objects.get(username='diego'), action='=HYPERLINK("http://x")')
        Audit.objects.create(user=None, action='Evento del sistema')
        self.login('lucia')
        everything = self.client.get('/audit.csv')
        self.assertEqual(everything.status_code, 200)
        self.assertIn('text/csv', everything['Content-Type'])
        text = everything.content.decode('utf-8-sig')
        self.assertIn("'=HYPERLINK", text)
        self.assertIn('Sistema', text)
        self.assertTrue(text.startswith('Fecha,Usuario,Evento'))
        filtered = self.client.get('/audit.csv', {'q': 'sistema'}).content.decode('utf-8-sig')
        self.assertIn('Evento del sistema', filtered)
        self.assertNotIn('HYPERLINK', filtered)

    def test_csv_requires_permission(self):
        self.login()
        self.assertEqual(self.client.get('/audit.csv').status_code, 403)


class WebRestore(BaseCase):
    def backup(self):
        self.login('admin')
        return self.client.post('/backup/').content

    def restore(self, content, confirm=True):
        data = {'file': SimpleUploadedFile('b.json', content)}
        if confirm:
            data['confirmar'] = '1'
        return self.client.post('/restore/', data, follow=True)

    def test_only_superusers_can_restore(self):
        content = self.backup()
        self.login('lucia')
        self.assertEqual(self.restore(content).status_code, 403)
        self.client.logout()
        self.assertEqual(self.restore(content).status_code, 403)
        self.login('lucia')
        self.assertNotContains(self.client.get('/p/backups/'), 'Restaurar un respaldo')
        self.login('admin')
        self.assertContains(self.client.get('/p/backups/'), 'Restaurar un respaldo')

    def test_requires_confirmation_and_a_valid_file(self):
        content = self.backup()
        self.assertContains(self.restore(content, confirm=False), 'casilla de confirmación')
        self.assertContains(self.restore(b'esto no es json'), 'no es un respaldo válido')
        self.assertContains(self.restore(b'{"a": 1}'), 'no es un respaldo válido')
        self.assertContains(self.client.post('/restore/', {'confirmar': '1'}, follow=True), 'Selecciona un archivo')

    def test_roundtrip_recovers_deleted_content_and_keeps_existing_users(self):
        content = self.backup()
        resource = Resource.objects.get(title='Documentación de Django')
        lesson = Lesson.objects.first()
        resource_id, lesson_id = resource.pk, lesson.pk
        ResourceVersion.objects.filter(resource=resource).delete()
        resource.delete()
        lesson.delete()
        response = self.restore(content)
        self.assertContains(response, 'Respaldo restaurado')
        self.assertTrue(Resource.objects.filter(pk=resource_id, title='Documentación de Django').exists())
        self.assertTrue(Lesson.objects.filter(pk=lesson_id).exists())
        self.assertTrue(self.client.login(username='ana', password=PASSWORD))
        self.assertTrue(Audit.objects.filter(action__startswith='Restauración de respaldo').exists())

    def test_new_users_come_back_without_a_password_and_with_their_groups(self):
        ghost = User.objects.create_user('fantasma', password='Clave-Que-No-Debe-Volver-1', first_name='Fanny')
        ghost.groups.add(Group.objects.get(name='Evaluador'))
        content = self.backup()
        ghost.delete()
        self.restore(content)
        restored = User.objects.get(username='fantasma')
        self.assertFalse(restored.has_usable_password())
        self.assertFalse(restored.check_password('Clave-Que-No-Debe-Volver-1'))
        self.assertEqual(restored.first_name, 'Fanny')
        self.assertEqual([g.name for g in restored.groups.all()], ['Evaluador'])
        self.assertNotIn(b'Clave-Que-No-Debe-Volver-1', content)

    def test_existing_user_data_is_not_overwritten(self):
        content = self.backup()
        User.objects.filter(username='ana').update(first_name='Cambiada')
        self.restore(content)
        self.assertEqual(User.objects.get(username='ana').first_name, 'Cambiada')

    def test_inconsistent_data_changes_nothing(self):
        content = json.loads(self.backup())
        content.append(
            {
                'model': 'hub.favorite',
                'pk': 999,
                'fields': {'user': 9999, 'resource': 9999, 'created': '2026-01-01T00:00:00Z'},
            }
        )
        before = Resource.objects.count()
        Resource.objects.filter(title='Documentación de Django').update(title='Editado')
        response = self.restore(json.dumps(content).encode())
        self.assertContains(response, 'No se cambió nada')
        self.assertEqual(Resource.objects.count(), before)
        self.assertTrue(Resource.objects.filter(title='Editado').exists())

    def test_unknown_models_are_skipped(self):
        content = json.loads(self.backup())
        content.append(
            {
                'model': 'sessions.session',
                'pk': 'abc',
                'fields': {'session_data': 'x', 'expire_date': '2030-01-01T00:00:00Z'},
            }
        )
        self.assertContains(self.restore(json.dumps(content).encode()), 'Respaldo restaurado')
        self.assertFalse(Session.objects.filter(pk='abc').exists())


class AssessmentNotices(BaseCase):
    def setUp(self):
        super().setUp()
        User.objects.filter(username__in=['ana', 'lucia']).update(email='x@example.com')
        User.objects.filter(username='ana').update(email='ana@example.com')
        User.objects.filter(username='lucia').update(email='lucia@example.com')

    def test_assigned_notice_goes_once_to_group_members_only(self):
        assessment = Assessment.objects.first()
        assessment.groups.add(Group.objects.get(name='Participante'))
        call_command('notify_assessments')
        recipients = sorted(m.to[0] for m in mail.outbox)
        self.assertEqual(recipients, ['ana@example.com'])
        self.assertIn('Evaluación asignada', mail.outbox[0].subject)
        call_command('notify_assessments')
        self.assertEqual(len(mail.outbox), 1)

    def test_no_assignment_notice_without_groups_or_before_opening(self):
        call_command('notify_assessments')
        self.assertEqual(len(mail.outbox), 0)
        assessment = Assessment.objects.first()
        assessment.groups.add(Group.objects.get(name='Participante'))
        assessment.opens_at = timezone.now() + timedelta(days=2)
        assessment.save()
        call_command('notify_assessments')
        self.assertEqual(len(mail.outbox), 0)

    def test_closing_reminder_skips_people_who_already_finished(self):
        assessment = Assessment.objects.first()
        assessment.closes_at = timezone.now() + timedelta(hours=20)
        assessment.save()
        ana = User.objects.get(username='ana')
        now = timezone.now()
        Attempt.objects.create(user=ana, assessment=assessment, deadline=now, finished=True, finished_at=now, score=50)
        call_command('notify_assessments')
        subjects = [(m.to[0], m.subject) for m in mail.outbox]
        self.assertIn(('lucia@example.com', f'Cierra pronto: {assessment.title}'), subjects)
        self.assertNotIn('ana@example.com', [to for to, _ in subjects])
        call_command('notify_assessments')
        self.assertEqual(len(mail.outbox), len(subjects))

    def test_far_deadlines_do_not_trigger_reminders(self):
        assessment = Assessment.objects.first()
        assessment.closes_at = timezone.now() + timedelta(days=10)
        assessment.save()
        call_command('notify_assessments')
        self.assertEqual(len(mail.outbox), 0)

    def test_assessments_without_questions_are_ignored(self):
        empty = Assessment.objects.create(
            title='Vacía', area='Transversal', closes_at=timezone.now() + timedelta(hours=1)
        )
        call_command('notify_assessments')
        self.assertFalse(Notification.objects.filter(assessment=empty).exists())
