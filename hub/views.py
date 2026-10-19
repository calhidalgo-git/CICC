import json
from datetime import timedelta
from io import BytesIO
from django.contrib.auth.decorators import login_required,permission_required
from django.contrib.auth.models import User,Group
from django.core import serializers
from django.db import transaction
from django.http import JsonResponse,HttpResponse,FileResponse,Http404
from django.shortcuts import render,redirect,get_object_or_404
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook,load_workbook
from .models import *
from .catalog import MODULES
QA_WEEKS = [
    'Semana 1 – Introducción al QA',
    'Semana 2 – Pruebas Manuales',
    'Semana 3 – Herramientas de QA',
    'Semana 4 – Automatización inicial',
    'Semana 5 – QA en metodologías ágiles y tradicionales',
    'Semana 6 – Tendencias y futuro del QA',
]
QA_ADVANCED_WEEKS = [
    'Semana 1 – Estrategias de pruebas avanzadas',
    'Semana 2 – Automatización intermedia',
    'Semana 3 – QA en CI/CD y DevOps',
    'Semana 4 – Testing basado en datos e IA',
    'Semana 5 – QA en metodologías híbridas',
    'Semana 6 – QA en pruebas de seguridad',
    'Semana 7 – QA en pruebas de rendimiento',
    'Semana 8 – QA en pruebas móviles',
    'Semana 9 – QA en pruebas de APIs y microservicios',
    'Semana 10 – QA en pruebas de accesibilidad y usabilidad',
    'Semana 11 – Soft skills y rol estratégico del QA',
]

@login_required
def qa_advanced(request):
    return render(request,'qa_advanced.html',ctx(request,'library',weeks=[title.split(' \u2013 ',1)[-1] for title in QA_ADVANCED_WEEKS]))

@login_required
def qa_advanced_week(request,week):
    if not 1 <= week <= len(QA_ADVANCED_WEEKS):
        raise Http404
    return render(request,'qa_advanced_week.html',ctx(request,'library',week=week,week_title=QA_ADVANCED_WEEKS[week-1],weeks=[title.split(' \u2013 ',1)[-1] for title in QA_ADVANCED_WEEKS],week_template=f'qa_advanced_weeks/week_{week}.html'))

@login_required
def qa_fundamentals(request):
    return render(request,'fundamentals.html',ctx(request,'library',weeks=[title.split(' \u2013 ',1)[-1] for title in QA_WEEKS]))

from django.db.models import Q
NAV=[('dashboard','Vista general','◈'),('library','Biblioteca','▤'),('videos','Videoteca','▷'),('links','Enlaces útiles','↗'),('paths','Rutas de aprendizaje','⌘'),('lessons','Mis lecciones','▧'),('assessments','Evaluaciones','◇'),('results','Mis resultados','✓'),('questions','Banco de preguntas','⊞'),('reports','Reportes','▥'),('users','Usuarios','♧'),('roles','Roles y permisos','⚿'),('settings','Configuración','⚙'),('audit','Auditoría','≡'),('backups','Respaldos','↺')]
NAV[1:1]=[(s,t,'▤') for s,t in MODULES.items()] + [('components','Mapa de componentes','⌘'),('search','Búsqueda global','⌕')]
GATES={'questions':'hub.change_question','reports':'hub.view_audit','users':'auth.view_user','roles':'auth.view_group','settings':'auth.change_group','audit':'hub.view_audit','backups':'hub.view_audit'}
def log(user,action): Audit.objects.create(user=user,action=action)
def ctx(request,section,**extra):
    return dict(nav=[(s,t,i) for s,t,i in NAV if s not in GATES or request.user.has_perm(GATES[s])],section=section,title=dict((s,t) for s,t,i in NAV).get(section,'Recurso'),**extra)
@login_required
def page(request,section='dashboard'):
    if section not in dict((s,t) for s,t,i in NAV): raise Http404
    if section in GATES and not request.user.has_perm(GATES[section]): return HttpResponse('Acceso restringido por tu rol',status=403)
    resources=Resource.objects.filter(published=True)
    query=request.GET.get('q',''); area=request.GET.get('area','')
    if query: resources=resources.filter(Q(title__icontains=query)|Q(description__icontains=query)|Q(area__icontains=query))
    if area: resources=resources.filter(area=area)
    if section in MODULES: resources=resources.filter(kind=section)
    if section=='videos': resources=resources.filter(kind='Video')
    if section=='links': resources=resources.filter(kind='Enlace')
    return render(request,'page.html',ctx(request,section,resources=resources,module_section=section in MODULES,search_lessons=Lesson.objects.filter(Q(title__icontains=query)|Q(body__icontains=query)|Q(path__title__icontains=query)) if query and section=='search' else Lesson.objects.none(),search_paths=LearningPath.objects.filter(Q(title__icontains=query)|Q(description__icontains=query)) if query and section=='search' else LearningPath.objects.none(),search_assessments=Assessment.objects.filter(Q(title__icontains=query)|Q(area__icontains=query)) if query and section=='search' else Assessment.objects.none(),chart_data=[Resource.objects.filter(area='Mainframe',published=True).count(),Resource.objects.filter(area='Distribuido',published=True).count()],query=query,area=area,paths=LearningPath.objects.all(),lessons=Lesson.objects.select_related('path'),done=list(Progress.objects.filter(user=request.user).values_list('lesson_id',flat=True)),assessments=Assessment.objects.all(),attempts=Attempt.objects.filter(user=request.user,finished=True).select_related('assessment'),questions=Question.objects.select_related('assessment'),users=User.objects.all(),roles=Group.objects.prefetch_related('permissions'),audits=Audit.objects.select_related('user').order_by('-created')[:100],resource_count=Resource.objects.count(),progress_count=Progress.objects.filter(user=request.user).count(),assessment_count=Assessment.objects.count()))
@login_required
def resource(request,pk):
    r=get_object_or_404(Resource,pk=pk,published=True)
    if r.title in ('Guía de pruebas batch y JCL', 'Fundamentos QA'):
        return render(request,'fundamentals.html',ctx(request,'library',weeks=[title.split(' \u2013 ',1)[-1] for title in QA_WEEKS]))
    if r.title in ('Checklist de pruebas de API REST', 'Checklist de pruebas API REST', 'QA Avanzado y Especialización'):
        return render(request,'qa_advanced.html',ctx(request,'library',weeks=[title.split(' \u2013 ',1)[-1] for title in QA_ADVANCED_WEEKS]))
    return render(request,'resource.html',ctx(request,'library',resource=r))
@login_required
def qa_week(request,week):
    if not 1 <= week <= len(QA_WEEKS):
        raise Http404
    return render(request,'qa_week.html',ctx(request,'library',week=week,week_title=QA_WEEKS[week-1],weeks=[title.split(' \u2013 ',1)[-1] for title in QA_WEEKS],week_template=f'qa_weeks/week_{week}.html'))

@login_required
def download(request,pk):
    v=get_object_or_404(ResourceVersion,pk=pk,resource__published=True)
    if not v.file: raise Http404
    log(request.user,f'Descarga recurso {v.resource_id} versión {v.number}')
    return FileResponse(v.file.open('rb'),as_attachment=True)
@login_required
@require_POST
def start(request,pk):
    a=get_object_or_404(Assessment,pk=pk)
    if not a.questions.exists(): return HttpResponse('Evaluación sin preguntas',status=400)
    active=Attempt.objects.filter(user=request.user,assessment=a,finished=False).first()
    if not active:
        active=Attempt.objects.create(user=request.user,assessment=a,deadline=timezone.now()+timedelta(minutes=a.minutes))
        log(request.user,f'Inicio evaluación {a.pk}')
    return redirect('exam',pk=active.pk)
@login_required
def exam(request,pk):
    a=get_object_or_404(Attempt,pk=pk,user=request.user)
    if a.finished: return redirect('results',pk=a.pk)
    return render(request,'exam.html',ctx(request,'assessments',attempt=a,exam_questions=a.assessment.questions.all(),deadline=a.deadline.isoformat()))
def grade(a):
    questions=list(a.assessment.questions.all())
    a.score=round(100*sum(a.answers.get(str(q.pk))==q.correct for q in questions)/len(questions),1) if questions else 0
    a.finished=True; a.save()
@login_required
@require_POST
@transaction.atomic
def save(request,pk):
    a=get_object_or_404(Attempt.objects.select_for_update(),pk=pk,user=request.user)
    if a.finished: return JsonResponse({'error':'Intento cerrado'},status=409)
    if timezone.now()>=a.deadline:
        grade(a); return JsonResponse({'expired':True},status=409)
    try:
        data=json.loads(request.body)
        qs={str(q.pk):q for q in a.assessment.questions.all()}
        if not isinstance(data,dict) or any(k not in qs or type(v)!=int or not 0<=v<len(qs[k].options) for k,v in data.items()): raise ValueError
    except (ValueError,TypeError): return JsonResponse({'error':'Respuestas inválidas'},status=400)
    a.answers=data; a.save(); return JsonResponse({'saved':True,'at':timezone.now().isoformat()})
@login_required
@require_POST
@transaction.atomic
def finish(request,pk):
    a=get_object_or_404(Attempt.objects.select_for_update(),pk=pk,user=request.user)
    if not a.finished: grade(a); log(request.user,f'Finaliza evaluación {a.pk}: {a.score}%')
    return redirect('results',pk=a.pk)
@login_required
def results(request,pk):
    a=get_object_or_404(Attempt,pk=pk,user=request.user,finished=True)
    rows=[(q,a.answers.get(str(q.pk)),q.options[q.correct]) for q in a.assessment.questions.all()]
    return render(request,'result.html',ctx(request,'results',attempt=a,rows=rows))
@login_required
@require_POST
def complete(request,pk):
    lesson=get_object_or_404(Lesson,pk=pk)
    Progress.objects.get_or_create(user=request.user,lesson=lesson)
    return redirect('/p/lessons/#lesson-'+str(pk))
def workbook_response(wb,name):
    out=BytesIO(); wb.save(out)
    response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']=f'attachment; filename="{name}"'; return response
@permission_required('hub.add_question',raise_exception=True)
def excel_template(request):
    wb=Workbook(); ws=wb.active; ws.append(['evaluacion_id','pregunta','opcion_a','opcion_b','opcion_c','correcta_0_2','explicacion']); ws.append([1,'¿Qué valida un smoke test?','Flujo crítico','Color del logo','Todos los casos',0,'Verifica viabilidad básica.'])
    return workbook_response(wb,'plantilla-preguntas.xlsx')
@permission_required('hub.add_question',raise_exception=True)
@require_POST
def import_excel(request):
    try:
        upload=request.FILES['file']
        if upload.size>2*1024*1024: raise ValueError('Máximo 2 MB')
        wb=load_workbook(upload,read_only=True,data_only=True)
        rows=list(wb.active.iter_rows(values_only=True))
        if len(rows)>501: raise ValueError('Máximo 500 preguntas')
        expected=('evaluacion_id','pregunta','opcion_a','opcion_b','opcion_c','correcta_0_2','explicacion')
        if tuple(rows[0])!=expected: raise ValueError('Encabezados incorrectos; usa la plantilla')
        with transaction.atomic():
            for n,row in enumerate(rows[1:],2):
                if len(row)!=7 or not all(row[i] is not None for i in range(7)) or type(row[5])!=int or row[5] not in range(3): raise ValueError(f'Fila {n}: datos inválidos')
                assessment=Assessment.objects.get(pk=int(row[0]))
                Question.objects.create(assessment=assessment,text=str(row[1]),options=list(map(str,row[2:5])),correct=row[5],explanation=str(row[6]))
            log(request.user,f'Importación Excel: {len(rows)-1} preguntas')
        return render(request,'notice.html',ctx(request,'questions',message=f'{len(rows)-1} preguntas importadas correctamente.'))
    except Exception as e:
        return render(request,'notice.html',ctx(request,'questions',message=f'No se importó ninguna fila. Revisa formato, campos e ID de evaluación. {str(e) if isinstance(e,ValueError) else "Archivo inválido"}'),status=400)
@permission_required('hub.view_audit',raise_exception=True)
def report(request):
    wb=Workbook(); ws=wb.active; ws.append(['Usuario','Evaluación','Área','Puntaje','Estado'])
    for a in Attempt.objects.filter(finished=True).select_related('user','assessment'):
        ws.append([a.user.username,a.assessment.title,a.assessment.area,a.score,'Aprobado' if a.score>=a.assessment.pass_score else 'Por reforzar'])
    log(request.user,'Exportación reporte Excel'); return workbook_response(wb,'reporte-cicc.xlsx')
@permission_required('hub.view_audit',raise_exception=True)
@require_POST
def backup(request):
    log(request.user,'Exportación respaldo de contenido')
    objects=[]
    for model in [Resource,ResourceVersion,LearningPath,Lesson,Assessment,Question,Progress,Attempt,Audit]: objects.extend(model.objects.all())
    response=HttpResponse(serializers.serialize('json',objects),content_type='application/json'); response['Content-Disposition']='attachment; filename="cicc-content-backup.json"'; return response
