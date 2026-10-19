"""Herramientas de gestión: plantilla e importación de preguntas, reporte Excel y respaldo."""

import csv
import json
import logging
from io import BytesIO
from zipfile import BadZipFile

from django.contrib import messages
from django.contrib.auth.decorators import permission_required
from django.contrib.auth.models import Group, User
from django.core import serializers
from django.core.exceptions import ObjectDoesNotExist, PermissionDenied, ValidationError
from django.core.serializers.base import DeserializationError
from django.core.serializers.json import DjangoJSONEncoder
from django.db import DatabaseError, connection, transaction
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..models import (
    Assessment,
    Attempt,
    Audit,
    Certificate,
    CourseWeek,
    Favorite,
    LearningPath,
    Lesson,
    LessonCheck,
    Notification,
    Progress,
    Question,
    Resource,
    ResourceVersion,
    WeekProgress,
)
from .common import ctx, filtered_attempts, filtered_audit, log

logger = logging.getLogger(__name__)

EXPECTED_HEADERS = ('evaluacion_id', 'pregunta', 'opcion_a', 'opcion_b', 'opcion_c', 'correcta_0_2', 'explicacion')
TOPIC_HEADER = 'tema'  # octava columna, opcional
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
MAX_REPORTED_ERRORS = 10
# Los campos marcados con = + - @ se interpretan como fórmulas en Excel; se neutralizan al exportar.
FORMULA_PREFIXES = ('=', '+', '-', '@')


class ImportRejected(Exception):
    """El archivo no cumple el formato. Lleva los mensajes que se muestran al usuario."""

    def __init__(self, *messages):
        super().__init__('; '.join(messages))
        self.messages = list(messages)


def workbook_response(wb, name):
    out = BytesIO()
    wb.save(out)
    response = HttpResponse(
        out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{name}"'
    return response


def safe_cell(value):
    return f"'{value}" if isinstance(value, str) and value.startswith(FORMULA_PREFIXES) else value


@permission_required('hub.add_question', raise_exception=True)
def excel_template(request):
    wb = Workbook()
    ws = wb.active
    ws.append([*EXPECTED_HEADERS, TOPIC_HEADER])
    ws.append(
        [
            1,
            '¿Qué valida un smoke test?',
            'Flujo crítico',
            'Color del logo',
            'Todos los casos',
            0,
            'Verifica viabilidad básica.',
            'Pruebas smoke',
        ]
    )
    return workbook_response(wb, 'plantilla-preguntas.xlsx')


def read_rows(upload):
    if upload.size > MAX_UPLOAD_BYTES:
        raise ImportRejected('El archivo supera el máximo de 2 MB.')
    try:
        wb = load_workbook(upload, read_only=True, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True)) if sheet else []
    except (BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise ImportRejected('El archivo no es un .xlsx válido.') from exc
    header = tuple(cell for cell in (rows[0] if rows else ()) if cell not in (None, ''))
    if header not in (EXPECTED_HEADERS, (*EXPECTED_HEADERS, TOPIC_HEADER)):
        raise ImportRejected('Encabezados incorrectos; usa la plantilla.')
    rows = [row for row in rows[1:] if any(cell not in (None, '') for cell in row)]
    if not rows:
        raise ImportRejected('El archivo no contiene preguntas.')
    if len(rows) > MAX_ROWS:
        raise ImportRejected(f'Máximo {MAX_ROWS} preguntas por archivo.')
    return rows


def build_question(number, row, assessments):
    """Valida una fila y devuelve la pregunta sin guardar; lanza ValueError con el motivo."""
    count = len(EXPECTED_HEADERS)
    required = row[:count]
    if len(required) != count or any(cell in (None, '') for cell in required):
        raise ValueError('hay celdas vacías')
    topic = str(row[count]).strip() if len(row) > count and row[count] else ''
    try:
        assessment = assessments[int(row[0])]
    except (KeyError, TypeError, ValueError):
        raise ValueError(f'la evaluación «{row[0]}» no existe') from None
    correct = row[5]
    if isinstance(correct, float) and correct.is_integer():
        correct = int(correct)
    if type(correct) is not int or correct not in range(3):
        raise ValueError('la columna correcta_0_2 debe ser 0, 1 o 2')
    question = Question(
        assessment=assessment,
        text=str(row[1]).strip(),
        options=[str(cell).strip() for cell in row[2:5]],
        correct=correct,
        explanation=str(row[6]).strip(),
        topic=topic,
    )
    try:
        question.full_clean()
    except ValidationError as exc:
        raise ValueError('; '.join(m for messages in exc.message_dict.values() for m in messages)) from exc
    return question


@permission_required('hub.add_question', raise_exception=True)
@require_POST
def import_excel(request):
    try:
        upload = request.FILES.get('file')
        if upload is None:
            raise ImportRejected('Selecciona un archivo .xlsx.')
        rows = read_rows(upload)
        assessments = Assessment.objects.in_bulk()
        questions, errors = [], []
        for number, row in enumerate(rows, 2):
            try:
                questions.append(build_question(number, row, assessments))
            except ValueError as exc:
                errors.append(f'Fila {number}: {exc}.')
        if errors:
            hidden = len(errors) - MAX_REPORTED_ERRORS
            raise ImportRejected(*errors[:MAX_REPORTED_ERRORS], *([f'… y {hidden} errores más.'] if hidden > 0 else []))
        with transaction.atomic():
            Question.objects.bulk_create(questions)
            log(request.user, f'Importación Excel: {len(questions)} preguntas')
    except ImportRejected as rejected:
        return render(
            request,
            'notice.html',
            ctx(request, 'questions', ok=False, message='No se importó ninguna fila.', details=rejected.messages),
            status=400,
        )
    except Exception:
        logger.exception('Error inesperado al importar preguntas')
        return render(
            request,
            'notice.html',
            ctx(
                request,
                'questions',
                ok=False,
                message='No se importó ninguna fila. Ocurrió un error inesperado.',
                details=[],
            ),
            status=500,
        )
    return render(
        request,
        'notice.html',
        ctx(request, 'questions', ok=True, message=f'{len(questions)} preguntas importadas correctamente.', details=[]),
    )


@permission_required('hub.view_audit', raise_exception=True)
def report(request):
    wb = Workbook()
    ws = wb.active
    ws.append(['Usuario', 'Evaluación', 'Área', 'Fecha', 'Puntaje', 'Estado'])
    for attempt in filtered_attempts(request.GET):
        ws.append(
            [
                safe_cell(attempt.user.username),
                safe_cell(attempt.assessment.title),
                attempt.assessment.area,
                attempt.started.strftime('%Y-%m-%d %H:%M'),
                attempt.score,
                'Aprobado' if attempt.score >= attempt.assessment.pass_score else 'Por reforzar',
            ]
        )
    log(request.user, 'Exportación reporte Excel')
    return workbook_response(wb, 'reporte-cicc.xlsx')


# Orden de respaldo y restauración: primero lo que otros objetos referencian.
BACKUP_MODELS = (
    Group,
    Resource,
    ResourceVersion,
    LearningPath,
    Lesson,
    LessonCheck,
    Assessment,
    Question,
    'users',
    Favorite,
    Progress,
    Attempt,
    Certificate,
    CourseWeek,
    WeekProgress,
    Notification,
    Audit,
)
RESTORABLE = {'auth.group', 'auth.user'} | {m._meta.label_lower for m in BACKUP_MODELS if not isinstance(m, str)}
MAX_RESTORE_BYTES = 50 * 1024 * 1024
USER_FIELDS = ('username', 'first_name', 'last_name', 'email', 'is_active', 'is_staff', 'is_superuser', 'groups')


@permission_required('hub.view_audit', raise_exception=True)
@require_POST
def backup(request):
    """Respaldo JSON restaurable con la pantalla de Respaldos o con ``manage.py loaddata``. Los usuarios van sin contraseña."""
    log(request.user, 'Exportación respaldo de contenido')
    objects = []
    for model in BACKUP_MODELS:
        if model == 'users':
            objects.extend(
                serializers.serialize('python', User.objects.all(), fields=USER_FIELDS, use_natural_foreign_keys=True)
            )
        else:
            objects.extend(serializers.serialize('python', model.objects.all(), use_natural_foreign_keys=True))
    response = HttpResponse(json.dumps(objects, cls=DjangoJSONEncoder), content_type='application/json')
    response['Content-Disposition'] = 'attachment; filename="cicc-content-backup.json"'
    return response


@require_POST
def restore(request):
    """Restaura un respaldo JSON. Solo superusuarios. Los objetos con el mismo ID se sobrescriben y los usuarios existentes se conservan."""
    if not request.user.is_authenticated or not request.user.is_superuser:
        raise PermissionDenied
    if not request.POST.get('confirmar'):
        messages.error(request, 'Marca la casilla de confirmación para restaurar.')
        return redirect('/p/backups/')
    upload = request.FILES.get('file')
    if upload is None or upload.size > MAX_RESTORE_BYTES:
        messages.error(request, 'Selecciona un archivo JSON de respaldo (máximo 50 MB).')
        return redirect('/p/backups/')
    try:
        objects = list(serializers.deserialize('json', upload.read().decode('utf-8')))
    except (UnicodeDecodeError, ValueError, DeserializationError):
        messages.error(request, 'El archivo no es un respaldo válido de CICC. No se cambió nada.')
        return redirect('/p/backups/')
    restored = skipped = 0
    try:
        with transaction.atomic():
            for item in objects:
                label = item.object._meta.label_lower
                if label not in RESTORABLE:
                    skipped += 1
                    continue
                if label == 'auth.user':
                    clash = User.objects.filter(Q(pk=item.object.pk) | Q(username=item.object.username)).exists()
                    if clash:
                        skipped += 1
                        continue
                    item.object.set_unusable_password()
                item.save()
                restored += 1
            # Con claves foráneas diferidas, un respaldo incoherente solo fallaría al confirmar: se comprueba aquí para revertir.
            connection.check_constraints()
    except (DatabaseError, ValueError, ObjectDoesNotExist):
        logger.exception('Error al restaurar el respaldo')
        messages.error(request, 'No se pudo restaurar el respaldo (datos inconsistentes). No se cambió nada.')
        return redirect('/p/backups/')
    log(request.user, f'Restauración de respaldo: {restored} objetos, {skipped} omitidos')
    messages.success(
        request,
        f'Respaldo restaurado: {restored} objetos. Se omitieron {skipped} (usuarios ya existentes u objetos no restaurables). '
        'Los usuarios nuevos necesitan una contraseña nueva.',
    )
    return redirect('/p/backups/')


@permission_required('hub.view_audit', raise_exception=True)
def audit_csv(request):
    """Exporta la auditoría (con los mismos filtros de la pantalla) a CSV para Excel."""
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="auditoria-cicc.csv"'
    response.write('\ufeff')
    writer = csv.writer(response)
    writer.writerow(['Fecha', 'Usuario', 'Evento'])
    for event in filtered_audit(request.GET)[:20000]:
        writer.writerow(
            [
                timezone.localtime(event.created).strftime('%Y-%m-%d %H:%M:%S'),
                safe_cell(event.user.username if event.user else 'Sistema'),
                safe_cell(event.action),
            ]
        )
    log(request.user, 'Exportación auditoría CSV')
    return response
